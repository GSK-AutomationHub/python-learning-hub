import openpyxl

string_list = [["column1", "column2", "column3"], 
               ["R1C1", "R1C2", "R1C3"],
               ["R2C1", "R1C2", "R1C3"]]

file_path = ".\\python_file_io\\test_excel.xlsx"

def write_to_excel(file_path):
    workbook = openpyxl.load_workbook(file_path)
    work_sheet = workbook.active
    rows = columns = 3
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            work_sheet.cell(row=row, column=column).value = string_list[row-1][column-1]
    workbook.save(file_path)

write_to_excel(file_path=file_path)

def read_the_excel(file_path):
    workbook = openpyxl.load_workbook(file_path)
    work_sheet = workbook.active
    rows = work_sheet.max_row
    columns = work_sheet.max_column
    for row in range(1, rows + 1):
        print("\n")
        for column in range(1, columns + 1):
            print(work_sheet.cell(row=row, column=column).value, end=" ")


read_the_excel(file_path=file_path)