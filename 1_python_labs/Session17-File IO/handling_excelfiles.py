import openpyxl
import os

mydict = {'header1':'sr.no','header2':'prog lan','header3':'book',
          'no1':1, 'lan1':'python','book1':'head-1st-python',
          'no2':2, 'lan2':'javascript','book2':'head-1st-js',
          'no3':3, 'lan3':'java','book3':'head-1st-java'}

file = os.getcwd() + '/sampleExcel.xlsx'


def get_workbook(file):
    workbook = openpyxl.load_workbook(file)
    return workbook

def get_active_sheet(workbook):
    active_sheet = workbook['Sheet1']
    return active_sheet

def get_max_rows(sheet):
    return sheet.max_row

def get_max_columns(sheet):
    return sheet.max_column

def write_to_excel(file):
    workbook = get_workbook(file)
    active_sheet = get_active_sheet(workbook)
    value_list = list(mydict.values())
    # print(value_list)
    rows = columns = 3
    i = 0
    for r in range(1, rows + 2):
        for c in range(1, columns + 1):
            active_sheet.cell(r, c, value_list[i])
            i += 1
    workbook.save(file)

def read_from_excel(file):
    workbook = get_workbook(file)
    active_sheet = get_active_sheet(workbook)
    rows = get_max_rows(active_sheet)
    columns = get_max_columns(active_sheet)
    for r in range (1, rows+1):
        for c in range(1, columns+1):
            print(active_sheet.cell(r,c).value,end=" ")
        print("\n")

def clear_the_excel(file):
    workbook = get_workbook(file)
    active_sheet = get_active_sheet(workbook)
    rows = get_max_rows(active_sheet)
    columns = get_max_columns(active_sheet)
    active_sheet.delete_rows(1)

write_to_excel(file)
read_from_excel(file)
clear_the_excel(file)
